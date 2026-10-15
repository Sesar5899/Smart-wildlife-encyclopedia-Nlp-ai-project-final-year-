from __future__ import annotations

import copy
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/myominko/Developer/animal")
REFERENCE = Path("/Users/myominko/Downloads/Final Body UniSync.docx")
OUTPUT = ROOT / "output" / "Smart Wildlife Encyclopedia using RAG and NLP.docx"
ASSETS = ROOT / "report-assets"

TITLE = "Smart Wildlife Encyclopedia using Retrieval-Augmented Generation & NLP"
PRODUCT = "Myanmar Wildlife Encyclopedia"
TNR = "Times New Roman"


def set_run_font(run, name=TNR, size=None, bold=None, italic=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_style_font(style, name=TNR, size=None, bold=None, italic=None):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def clear_body(document):
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def ensure_style(document, name, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def configure_styles(document):
    normal = document.styles["Normal"]
    set_style_font(normal, size=12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    body = document.styles["Body"]
    set_style_font(body, size=12)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.first_line_indent = Inches(0.5)
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(6)

    bullet = document.styles["List Bullet"]
    set_style_font(bullet, size=12)
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bullet.paragraph_format.line_spacing = 1.5
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(3)

    numbered = document.styles["List Number"]
    set_style_font(numbered, size=12)
    numbered.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    numbered.paragraph_format.line_spacing = 1.5
    numbered.paragraph_format.left_indent = Inches(0.5)
    numbered.paragraph_format.first_line_indent = Inches(-0.25)
    numbered.paragraph_format.space_after = Pt(3)

    chapter_label = ensure_style(document, "Report Chapter Label")
    set_style_font(chapter_label, size=14, bold=True)
    chapter_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter_label.paragraph_format.line_spacing = 1.5
    chapter_label.paragraph_format.space_before = Pt(6)
    chapter_label.paragraph_format.space_after = Pt(0)
    chapter_label.paragraph_format.keep_with_next = True

    chapter_title = ensure_style(document, "Report Chapter Title")
    set_style_font(chapter_title, size=14, bold=True)
    chapter_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter_title.paragraph_format.line_spacing = 1.5
    chapter_title.paragraph_format.space_before = Pt(0)
    chapter_title.paragraph_format.space_after = Pt(6)
    chapter_title.paragraph_format.keep_with_next = True

    section = ensure_style(document, "Report Section")
    set_style_font(section, size=14, bold=True)
    section.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    section.paragraph_format.line_spacing = 1.5
    section.paragraph_format.space_before = Pt(6)
    section.paragraph_format.space_after = Pt(6)
    section.paragraph_format.keep_with_next = True

    subsection = ensure_style(document, "Report Subsection")
    set_style_font(subsection, size=12, bold=True)
    subsection.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subsection.paragraph_format.line_spacing = 1.5
    subsection.paragraph_format.space_before = Pt(6)
    subsection.paragraph_format.space_after = Pt(3)
    subsection.paragraph_format.keep_with_next = True

    caption = ensure_style(document, "Report Figure Caption")
    set_style_font(caption, size=12)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)

    table_caption = ensure_style(document, "Report Table Caption")
    set_style_font(table_caption, size=12)
    table_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_caption.paragraph_format.line_spacing = 1.5
    table_caption.paragraph_format.space_before = Pt(6)
    table_caption.paragraph_format.space_after = Pt(6)
    table_caption.paragraph_format.keep_with_next = True

    code = ensure_style(document, "Report Command")
    set_style_font(code, name="Consolas", size=8.5)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.first_line_indent = Inches(0)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.keep_together = True

    refs = ensure_style(document, "Report Reference")
    set_style_font(refs, size=11)
    refs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    refs.paragraph_format.line_spacing = 1.5
    refs.paragraph_format.left_indent = Inches(0.35)
    refs.paragraph_format.first_line_indent = Inches(-0.35)
    refs.paragraph_format.space_after = Pt(3)


def body(document, text):
    p = document.add_paragraph(style="Body")
    p.add_run(text)
    return p


def bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def numbered(document, text):
    p = document.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def chapter(document, number, title, first=False):
    if not first:
        document.add_page_break()
    p = document.add_paragraph(style="Report Chapter Label")
    p.add_run(f"CHAPTER {number}")
    p = document.add_paragraph(style="Report Chapter Title")
    p.add_run(title.upper())


def section_heading(document, text):
    p = document.add_paragraph(style="Report Section")
    p.add_run(text)
    return p


def subsection_heading(document, text):
    p = document.add_paragraph(style="Report Subsection")
    p.add_run(text)
    return p


def figure(document, image_path, caption_text, width=5.77):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(image_path), width=Inches(width))
    picture._inline.docPr.set("descr", caption_text)
    picture._inline.docPr.set("title", caption_text)
    c = document.add_paragraph(style="Report Figure Caption")
    c.add_run(caption_text)
    return c


def command_block(document, text):
    p = document.add_paragraph(style="Report Command")
    run = p.add_run(text)
    set_run_font(run, name="Consolas", size=8.5)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "B7B7B7")
        borders.append(el)
    ppr.append(borders)
    return p


TABLE_WIDTHS = [1710, 1350, 2625, 3345]


def set_cell_text(cell, text, bold=False, size=7.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=55, start=75, bottom=55, end=75):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def data_table(document, caption_text, rows):
    cap = document.add_paragraph(style="Report Table Caption")
    cap.add_run(caption_text)
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Normal Table"
    headers = ["Attribute\nName", "Data Type", "Key Constraint", "Description"]
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_text(cell, header, bold=True, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.width = Inches(TABLE_WIDTHS[index] / 1440)
        set_cell_margins(cell)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "D9E2F3")
        cell._tc.get_or_add_tcPr().append(shd)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)

    for row_values in rows:
        row = table.add_row()
        for index, (cell, value) in enumerate(zip(row.cells, row_values)):
            set_cell_text(cell, value, size=7.3)
            cell.width = Inches(TABLE_WIDTHS[index] / 1440)
            set_cell_margins(cell)

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(TABLE_WIDTHS)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in TABLE_WIDTHS:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(TABLE_WIDTHS[index]))
            tc_w.set(qn("w:type"), "dxa")
    return table


def enable_field_updates(document):
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def add_reference(document, number, text):
    p = document.add_paragraph(style="Report Reference")
    p.add_run(f"[{number}] {text}")


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(REFERENCE)
    clear_body(document)
    configure_styles(document)
    enable_field_updates(document)
    document.core_properties.title = TITLE
    document.core_properties.subject = "Final project report"
    document.core_properties.keywords = "wildlife, encyclopedia, RAG, NLP, Myanmar, Flask, SQLite"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""

    # Reassert the user-requested page system even though it is inherited from the reference.
    for s in document.sections:
        s.page_width = Inches(8.26875)
        s.page_height = Inches(11.69375)
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        s.header_distance = Inches(0.5)
        s.footer_distance = Inches(0.5)

    chapter(document, 1, "Introduction", first=True)
    body(document, f"Biodiversity information is distributed across scientific catalogues, legal schedules, conservation assessments, photographs, and narrative records. For a learner or member of the public, the difficulty is not simply finding a web page; it is finding a trustworthy description, understanding technical conservation terms, locating the same animal under Myanmar, English, and scientific names, and asking a follow-up question without receiving an invented answer. The project entitled {TITLE} addresses this problem through a source-centred bilingual web encyclopedia. Its deployed product identity is {PRODUCT}, and its initial collection focuses on Myanmar species, ecosystems, and protected areas.")
    body(document, "The application is implemented as a Flask and SQLite web system with server-rendered Jinja interfaces, local JSON catalogues, verified WebP photographs, Myanmar-aware search, account-controlled access, and a species-specific conversational workspace. At the time of this report the local catalogue contains 182 animal or protected-taxon records in three categories, 64 ecosystem records, and 73 protected-area records. The figures shown by the interface are read from the actual SQLite and JSON data rather than being typed into a static landing page. This makes the public summary consistent with the records available to the member.")
    body(document, "Retrieval-Augmented Generation (RAG) is used to constrain conversational answers to the selected record. Each animal record is divided into labelled passages such as identity, habitat, diet, population, threats, legal protection, and source information. A Myanmar-aware lexical retriever chooses relevant passages, while a local cache of OpenAI embedding vectors can add semantically similar passages when an administrator has already indexed the material. The selected evidence is placed in a strict system prompt and sent to a Groq-hosted OpenAI-compatible chat model. The model is not called when the request is outside the record, asks for another subject, or asks for a photograph that can be resolved deterministically from local files.")
    body(document, "This report explains the motivation, theoretical background, architecture, permanent storage, implementation, installation, verification, limitations, and future direction of the project. The discussion is grounded in the current repository and database. System diagrams are derived from the implemented routes, modules, gates, tables, and external interfaces, and the implementation chapter uses desktop screenshots captured from the running application.")

    section_heading(document, "1.1\tMotivation")
    body(document, "A conventional wildlife website can provide long pages of facts, but the information is often difficult to search in Myanmar script and difficult to question conversationally. A general chatbot can answer quickly, yet it may blend unrelated knowledge, state uncertain information confidently, or cite a page that did not support the answer. These limitations are important in biodiversity communication because species names, legal protection, conservation status, and photographic identity must not be guessed.")
    bullet(document, "Fragmented evidence: species facts, ecosystem descriptions, protected-area records, legal schedules, conservation assessments, and photographs are usually stored in different places.")
    bullet(document, "Myanmar text variation: invisible characters, alternative mark order, transliteration, aliases, and head nouns can prevent a literal database lookup from finding the intended species.")
    bullet(document, "Hallucination risk: a general language model may answer beyond the selected wildlife record unless retrieval, scope checks, and refusal rules are enforced before generation.")
    bullet(document, "Subject confusion: a question typed in one species chat may name another species, an ecosystem, or a protected area, and the application must hand the user to the correct record rather than answering in the wrong context.")
    bullet(document, "Media integrity: a missing female, male, or cover photograph must be reported as missing; substituting another animal or an unlabeled image would create false evidence.")
    bullet(document, "Provenance and licensing: imported photographs and descriptions require source URLs, credits, and ownership information to remain attached to the published record.")
    bullet(document, "Administrative transparency: the system maintainer needs to see why a passage was or was not retrieved and whether a weak answer is a retrieval problem or a generator problem.")
    body(document, "The project is therefore motivated by two objectives that must be satisfied together: accessible wildlife learning and engineering honesty. Accessibility requires bilingual navigation, forgiving search, attractive record pages, and conversational follow-up. Honesty requires local evidence, deterministic gates, source attribution, protected account data, explicit missing-information responses, and an inspectable retrieval pipeline.")

    section_heading(document, "1.2\tAim and Objectives")
    body(document, f"The principal aim of {TITLE} is to provide a searchable, source-backed, Myanmar-oriented wildlife encyclopedia in which a signed-in member can browse reliable records and ask questions that are answered only from the selected species evidence. The system also gives an administrator controlled publishing and retrieval-diagnostic tools.")
    bullet(document, "Publish a unified catalogue of Myanmar animal records with Myanmar, English, and scientific names, taxonomy, habitat, diet, population, conservation, legal-protection, source, and photograph fields.")
    bullet(document, "Provide separate catalogues for ecosystems and protected areas with filters, pagination, bilingual labels, local photographs, and source links.")
    bullet(document, "Require an account for member content, return members to their originally requested page after login, and prevent the login redirection parameter from leaving the site.")
    bullet(document, "Implement Myanmar-aware search that normalizes Unicode, tolerates character-order variation, uses aliases and phonetic similarity where safe, and falls back to SQLite FTS5 for descriptive text.")
    bullet(document, "Create one conversational workspace per species rather than a global chat, so the active subject and evidence boundary are always explicit.")
    bullet(document, "Retrieve labelled passages with lexical scoring and intent bonuses, allow cached semantic vectors to fill unused evidence slots, and keep lexical hits in authoritative order.")
    bullet(document, "Refuse out-of-scope questions before generation, resolve photo requests from local media, and redirect questions about another record with one clear navigation action.")
    bullet(document, "Stream grounded answers from the server, display the sources used, and save completed exchanges privately for the signed-in member and selected species.")
    bullet(document, "Provide Content Studio for authorized administrators to create, edit, delete, and photograph records with CSRF and upload validation.")
    bullet(document, "Provide a twelve-stage retrieval console and an administrator-only embedding-index operation for all three content catalogues.")
    bullet(document, "Verify core behavior through automated tests covering data integrity, access control, search, retrieval, photographs, streaming, history privacy, administrative publishing, and vector persistence.")

    section_heading(document, "1.3\tAdvantage of the System")
    body(document, "The project combines functions that are commonly separated among catalogue sites, search engines, generic chatbots, local media folders, and administrative spreadsheets. Its main advantage is not the presence of an AI model alone, but the controlled path from curated record to retrieved passage to visible answer.")
    bullet(document, "Source-bounded answers: only passages selected from the active record are sent to the generator, and a no-match decision produces a refusal instead of a speculative completion.")
    bullet(document, "Myanmar-aware discovery: normalization, aliases, head-noun logic, phonetic comparison, fuzzy scores, and FTS5 improve access to records despite real input variation.")
    bullet(document, "Deterministic safety branches: different-subject requests, photograph requests, and unsupported questions are handled by application code with predictable wording and behavior.")
    bullet(document, "Hybrid retrieval without hidden online dependency: lexical retrieval works locally, stored vectors may add semantic recall, and member requests never create OpenAI embeddings.")
    bullet(document, "Transparent administration: the retrieval console exposes intermediate units, intent detection, passage scores, threshold decisions, prompt context, and final response branch.")
    bullet(document, "Data provenance: source URLs, portal identifiers, image credits, legal bases, and crawl timestamps stay with the record and are available to the interface and answer sources.")
    bullet(document, "Privacy boundaries: chat history belongs to one user and one species; browser-supplied history is ignored in favour of trusted server history.")
    bullet(document, "Maintainable local deployment: Flask modules, Jinja templates, SQLite, local JSON snapshots, and static WebP assets create a complete desktop web application that can be inspected and tested without a separate frontend build.")
    body(document, "These advantages make the system suitable as an educational and conservation-information prototype. It remains intentionally cautious: it reports missing fields, separates IUCN threat categories from Myanmar legal protection, does not pretend a taxonomic family is a species, and uses image credits where the source licence requires them.")

    section_heading(document, "1.4\tSummary")
    body(document, f"This chapter introduces {TITLE} as a bilingual, source-centred wildlife information system. The project responds to fragmented biodiversity evidence, Myanmar search difficulty, hallucination risk, photograph misidentification, and the need for administrative traceability. Its objectives cover catalogues, authentication, Myanmar-aware search, species-specific RAG, deterministic refusal and media handling, private history, publishing, indexing, and verification. The result is an application in which learning convenience is constrained by explicit evidence and provenance rules.")

    chapter(document, 2, "Background Theory")
    body(document, "This chapter presents the theories and technologies that support the wildlife encyclopedia. The implementation joins biodiversity informatics, natural-language processing, full-text and fuzzy search, retrieval-augmented generation, vector embeddings, large-language-model streaming, access control, and data provenance. Each subject is described in relation to the decisions visible in the current system rather than as an isolated technology survey.")

    section_heading(document, "2.1\tTheoretical Foundation")
    body(document, "A wildlife record is a structured description of an entity, but a user's question is unstructured language. The central theoretical problem is therefore evidence alignment: identifying the record, selecting the smallest passages that answer the question, and producing a readable response that does not cross the evidence boundary. The project handles this problem through several cooperating layers.")

    subsection_heading(document, "2.1.1\tBiodiversity Informatics and Digital Encyclopedias")
    body(document, "Biodiversity informatics organizes names, taxonomic ranks, distributions, habitats, populations, threats, conservation assessments, legal protection, media, and sources as data that can be searched and compared. A digital encyclopedia adds an explanatory interface over that data. In this project, animals are relational records because their categories, aliases, account histories, and vector mappings benefit from explicit keys. Ecosystems and protected areas are complete local JSON snapshots because they are reference catalogues with rich and heterogeneous assessment fields.")
    body(document, "Taxonomic identity requires more than a display name. A record can have a Myanmar name, English name, scientific name, aliases, order, family, linked portal name, and rank. The system preserves these distinctions so that a search can find the record without rewriting its scientific identity. Legal protection is stored separately from global or national conservation status because they answer different questions: extinction risk is an assessment, while protection status is a rule under a legal schedule.")

    subsection_heading(document, "2.1.2\tMyanmar Natural-Language Processing")
    body(document, "Myanmar script creates retrieval challenges that are not solved by lower-casing or whitespace tokenization alone. Visually identical text can contain invisible join controls or different Unicode mark order. A user may type a common head noun rather than the full species name, and phonetic spellings can vary. The repository therefore contains a small reusable Myanmar normalization and phonetic package. Search first canonicalizes the input, builds comparable keys, evaluates literal and alias matches, and uses phonetic equivalence only where it does not collapse common grammar into a false species name.")
    body(document, "The RAG retriever uses Myanmar syllable runs and adjacent pairs as strong lexical units while keeping whole Latin words. Stopword removal reduces weak function words, and intent detection recognizes concepts such as habitat, diet, population, threats, photographs, sources, and legal status. This design is corpus-aware: it values the forms that appear in the stored Myanmar wildlife records instead of assuming that an English tokenizer will divide the language correctly.")

    subsection_heading(document, "2.1.3\tFull-Text, Fuzzy, and Entity Search")
    body(document, "Entity search and evidence retrieval have different goals. Entity search must identify which record the user means; evidence retrieval must identify which fields inside that record answer a question. The catalogue search combines exact normalized names, aliases, phonetic similarity, fuzzy scores, and an SQLite FTS5 virtual table over Myanmar name, English name, scientific name, and description. FTS5 is an SQLite virtual-table module designed for efficient full-text queries over large text collections [3]. Triggers keep the external-content index synchronized with inserts, updates, and deletes in the animals table.")
    body(document, "A fixed similarity threshold avoids returning every weak candidate. Head-noun matching prefers the longest unambiguous name and deliberately returns no record when one short term matches many species. This prevents a generic term such as the Myanmar word for turtle from selecting one arbitrary turtle. Search therefore balances recall with entity safety, and the displayed results retain a match source and score that can be tested.")

    subsection_heading(document, "2.1.4\tRetrieval-Augmented Generation")
    body(document, "Retrieval-Augmented Generation combines a language generator with external non-parametric memory. The original RAG formulation was proposed to improve knowledge-intensive generation by retrieving explicit documents rather than relying only on facts stored in model parameters [2]. In this project the non-parametric memory is not an unrestricted web index. It is the current species record, transformed into labelled passages with stable field meanings.")
    body(document, "The retriever calculates how much of the question's strong lexical unit set overlaps each passage and adds a bonus when detected question intent matches the passage tags. Identity passages are appended for subject clarity but do not occupy semantic evidence slots. A matched result can contain several passages, but the generator receives only those selected passages and the record's identity. If no relevant passage clears the lexical guarantee and no permitted semantic addition applies, the system returns a deterministic not-available message.")

    subsection_heading(document, "2.1.5\tEmbeddings and Semantic Similarity")
    body(document, "An embedding represents text as a numeric vector whose geometry reflects aspects of semantic similarity. The project uses OpenAI's text-embedding-3-small model for administrator-created passage and query vectors. The default vector length for this model is 1,536 dimensions, which matches the local configuration [5]. Vectors are normalized and stored as binary arrays in SQLite, keyed by model name and the SHA-256 hash of the passage text.")
    body(document, "Cosine similarity compares the direction of two vectors. OpenAI recommends cosine similarity for its normalized embeddings, and cosine and Euclidean distance produce the same ranking for normalized vectors [5]. A crucial limitation is that every query has a nearest vector even when it is unrelated. The system therefore does not use semantic similarity as the refusal guarantee. Lexical logic decides whether the question is sufficiently grounded; semantic matches may only add passages to empty slots above the configured threshold and never displace or reorder lexical evidence.")

    subsection_heading(document, "2.1.6\tLarge Language Models, Prompt Grounding, and Streaming")
    body(document, "A large language model generates a sequence of tokens conditioned on a prompt. In a grounded question-answering system, prompt construction is as important as model choice. The system prompt names the current species, states that only the supplied passages may be used, requires acknowledgement of unavailable information, and provides source context separately from the member's question. Trusted chat history is loaded by the server; history supplied by a browser is ignored so a client cannot inject an authoritative earlier message.")
    body(document, "The generator is accessed through Groq's OpenAI-compatible API. Groq documents that its endpoint is mostly compatible with OpenAI client libraries and uses the base URL https://api.groq.com/openai/v1 [6]. The application opens a streamed completion and parses only data chunks into answer text. Server-sent events carry incremental deltas to the browser, followed by sources and a completion record. The exchange is saved only after completion so an interrupted answer does not become trusted history.")

    subsection_heading(document, "2.1.7\tDeterministic Gates and Honest Refusal")
    body(document, "The safest model call is sometimes no model call. Before retrieval and generation, the application evaluates three deterministic branches. The subject gate detects whether the member named another species, ecosystem, or protected area and returns one navigation action to that record. The photo gate interprets requests for cover, male, female, or all photographs and resolves only local static assets. The scope gate rejects a question when the active record contains no evidence for it. These branches prevent unnecessary external requests and make their responses repeatable.")
    body(document, "Honest refusal is not a failure of the application; it is a designed result that preserves the contract with the reader. A species page cannot legitimately answer a football score, financial price, recipe, or fact about a different animal. Similarly, a request for a female image cannot be satisfied by an unlabeled cover photograph. The refusal and missing-media messages make the evidence boundary visible.")

    subsection_heading(document, "2.1.8\tWeb Security, Accounts, and Role Separation")
    body(document, "Flask is a lightweight WSGI web framework that supplies routing, request processing, templates, sessions, and extension points [4]. The project uses signed sessions to associate browser requests with a user record and Werkzeug password hashing rather than storing plaintext passwords. Only the landing, registration, and login pages are public. Member pages redirect anonymous browser requests to login, while JSON APIs return HTTP 401 so client code does not receive an HTML form unexpectedly.")
    body(document, "Administrator access is explicit. An is_admin flag or configured administrator email is checked before Content Studio routes are served. Administrative mutations require a CSRF token, uploaded photographs are validated by content signature and size, and the application's maximum request size is configurable. The login next parameter accepts only same-site relative paths, preventing an open redirect. API keys remain server-side environment variables and are never rendered into templates or sent to browser JavaScript.")

    subsection_heading(document, "2.1.9\tProvenance, Local Media, and Data Quality")
    body(document, "Provenance records where a fact or image came from and supports later verification. The animal table retains portal record identifiers, source URLs, source names, image-source URLs, crawl time, image credit, and a JSON copy of portal material. Ecosystem and protected-area snapshots retain the same source relationship. Importers verify non-empty WebP files and deliberately publish a missing-photo panel when no acceptable image exists. Protected-species importers preserve legal bases and distinguish a species explicitly named in a schedule from one protected through its family, genus, or order.")
    body(document, "This approach avoids a common presentation error: showing a convenient but incorrect image. A group-level taxon may display an example member only when the credit clearly says that the image is an example from the group. Wikimedia photographs retain the credit required by their licences. These measures make media quality part of the data model rather than a purely visual concern.")

    section_heading(document, "2.2\tProposed Method")
    body(document, "The proposed method follows a source-to-answer pipeline. Each stage has a narrow responsibility and leaves an inspectable result that the administrator console can display. The same functions used by the member chat produce the trace, which prevents documentation code from describing a pipeline different from the one actually executed.")
    numbered(document, "Acquire and curate: import species, legal schedules, ecosystems, protected areas, and verified images into local SQLite, JSON, and static-media storage while retaining provenance.")
    numbered(document, "Identify the subject: normalize the query, compare names and aliases, reject ambiguous head nouns, and keep a species chat bound to one stable record key.")
    numbered(document, "Build evidence passages: convert present record fields into labelled passages with intent tags and identity flags; omit absent fields so suggestions cannot promise unavailable information.")
    numbered(document, "Apply deterministic gates: resolve other-record, photograph, and out-of-scope branches before any model request.")
    numbered(document, "Retrieve lexically: segment the question, remove weak units, detect intent, calculate overlap scores, compare the active threshold, and preserve record order.")
    numbered(document, "Add semantic recall: when an exact administrator-cached query vector exists, compare it with locally stored passage vectors and fill unused passage slots above the threshold.")
    numbered(document, "Generate from context: create a strict prompt from the selected passages and trusted history, open the configured Groq stream, and forward valid deltas as server-sent events.")
    numbered(document, "Persist and explain: save the completed exchange by user and animal key, display evidence and photographs, and expose the same branch and scores in the administrator retrieval console.")

    subsection_heading(document, "2.2.1\tData Acquisition and Publication")
    body(document, "The initial species catalogue is assembled from cached Myanmar Biodiversity Portal records and protected-wildlife schedules. Importers are designed to be repeatable: the portal crawl can rebuild from the local cache; legal importers add or refresh the rows they own without overwriting richer portal records; and backups are created before destructive database replacement. Ecosystems and protected areas are stored as complete local snapshots, allowing filters and details to work without an online portal request.")

    subsection_heading(document, "2.2.2\tIndexing and Retrieval Preparation")
    body(document, "SQLite initialization creates relational tables, indexes, FTS5 triggers, and forward-compatible columns. The administrator may run a full embedding build that visits animal, ecosystem, and protected-area records, hashes each passage, requests only missing vectors, maps vectors back to every record position, and removes stale mappings after a successful full run. The process is resumable because vectors are deduplicated by model and text hash.")

    subsection_heading(document, "2.2.3\tMember Question Processing")
    body(document, "A member opens a species chat and submits a message to the JSON endpoint. Authentication, input length, record existence, and subject checks run first. The server then evaluates deterministic response branches. For an answerable question it retrieves passages, constructs messages, streams the model response, saves the final exchange, and returns evidence metadata. The browser renders Markdown, sources, local images, and recent species conversations without trusting client-supplied history.")

    subsection_heading(document, "2.2.4\tAdministration and Evaluation")
    body(document, "Content Studio gives an administrator a controlled publishing surface for animals, ecosystems, and protected areas. Saving a record makes it immediately available to catalogues and species chat, and an online administrator can index the changed record. The retrieval console runs questions against any record and displays twelve stages including normalization, segmentation, stopword removal, strong and weak units, scope and photo gates, passage inventory, intents, per-passage scores, semantic columns, selected context, system prompt, and final response branch. A generated answer is offered as an explicit action rather than being run merely because the page opened.")

    section_heading(document, "2.3\tSummary")
    body(document, "This chapter establishes the foundation for a source-grounded wildlife encyclopedia. Biodiversity data modelling preserves taxonomic, legal, media, and provenance distinctions. Myanmar NLP and entity search address real input variation. FTS5, lexical overlap, intents, cached embeddings, and cosine similarity support hybrid retrieval, while deterministic subject, photo, and scope gates preserve an honest refusal boundary. Flask sessions, role separation, CSRF, upload checks, and server-side keys protect the web workflow. The proposed method links curated records to inspectable passages and streams only those answers that the local evidence permits.")

    chapter(document, 3, "System Architecture and Design")
    body(document, f"This chapter presents the architecture and design of {TITLE}. The system is organized as a layered desktop web application whose central rule is that a conversational answer must remain attached to one selected wildlife record. The design covers functional roles, request paths, retrieval logic, external AI boundaries, permanent storage, JSON catalogues, media, and administrator operations.")

    section_heading(document, "3.1\tSystem Architecture")
    body(document, "The architecture separates visible interaction from security checks, retrieval, external generation, and persistent evidence. A browser request first reaches a Flask route or JSON endpoint. Sessions and role checks decide whether the visitor is a member or administrator. Catalogue queries use SQLite or local JSON. A chat query passes through subject, photograph, and scope gates before retrieval. Only the matched generation branch reaches Groq, while OpenAI embeddings are created only through administrator operations. This separation reduces the number of components that can change a member answer.")
    numbered(document, "Client Layer: desktop and laptop browsers display the public landing page, account forms, member catalogues, species details, species chat, search, and administrator tools.")
    numbered(document, "Presentation Layer: Jinja2 templates render server data, the shared CSS system controls desktop and responsive layout, and JavaScript supports search suggestions, Markdown rendering, chat streaming, history deletion, and administrative interactions.")
    numbered(document, "Flask Web and Security Layer: page routes and JSON APIs enforce signed sessions, password verification, member access, administrator role checks, CSRF, safe redirects, input limits, and upload validation.")
    numbered(document, "Myanmar NLP and Retrieval Layer: normalization, syllable units, alias and subject resolution, intent detection, passage building, lexical scoring, semantic recall, prompt construction, and response-branch tracing are implemented in reusable Python modules.")
    numbered(document, "AI and Integration Layer: the generator uses Groq's OpenAI-compatible chat endpoint; the administrator indexer uses OpenAI embeddings; importers connect published biodiversity and legal sources to local records.")
    numbered(document, "Data and Media Layer: SQLite stores animals, users, history, FTS5 content, and vectors; JSON stores ecosystem and protected-area snapshots; static folders store verified WebP photographs.")
    numbered(document, "Published Sources and Delivery: portal records, protected-wildlife schedules, IUCN-related data sources, Wikidata, GBIF, Wikipedia, Wikimedia Commons, and the selected WSGI delivery environment supply or serve the system.")
    figure(document, ASSETS / "diagrams" / "system-architecture.png", f"Figure 3.1: System Architecture of {TITLE}", width=5.77)

    subsection_heading(document, "3.1.1\tFunctional Requirement and Use Case Diagram")
    body(document, "The system has three principal human roles. A visitor can inspect the public introduction and create or access an account. A member can browse all published catalogues and use species-specific assistance. An administrator has every member capability plus publishing, photograph management, diagnostic, and indexing functions. Groq and OpenAI are shown as external AI services because they are reached through server-side interfaces rather than being part of the browser.")
    bullet(document, "Visitor requirements: view the landing page, see live catalogue totals, register with validated account fields, sign in, and follow a safe return path.")
    bullet(document, "Member catalogue requirements: browse species by category; search Myanmar, English, scientific, alias, and description text; view detailed species evidence; and browse ecosystem and protected-area reference catalogues.")
    bullet(document, "Member chat requirements: open the chat for one species, ask a grounded question, receive deterministic media or refusal responses when appropriate, stream matched answers, see source metadata, restore recent conversations, and clear only the member's own history.")
    bullet(document, "Administrator requirements: open Content Studio, create or edit records, validate photographs, delete authorized records, inspect retrieval stages, compare generator behavior, and build or refresh the passage-vector index.")
    bullet(document, "System requirements: keep credentials and API keys on the server, preserve source attribution, prevent cross-account history access, and avoid model calls for unsupported questions.")
    figure(document, ASSETS / "diagrams" / "use-case-diagram.png", f"Figure 3.2: Use Case Diagram of {TITLE}", width=6.04)

    subsection_heading(document, "3.1.2\tSystem Components and Sequence Diagram")
    body(document, "The sequence diagram describes the highest-risk workflow: a member asks a question in a species chat. The browser sends only the current message and record identifier. The Flask API resolves trusted session state and the current record, then asks the gates whether the message names another subject, requests photographs, or lacks record evidence. A deterministic branch can immediately return a redirect, local image, or refusal without opening a generator stream.")
    body(document, "For an in-scope question, the RAG component builds labelled passages and reads lexical candidates, optional cached query vectors, and passage vectors. Myanmar units and question intents produce lexical scores; semantic recall may fill unused slots above its threshold. The selected passages and evidence mode return to the API. A grounded prompt is sent to Groq, answer deltas are forwarded to the browser, and the completed exchange is saved to chat history. The browser renders the answer, sources, and images. An invalid or anonymous API request ends before any of these retrieval or generation operations.")
    figure(document, ASSETS / "diagrams" / "sequence-diagram.png", f"Figure 3.3: Sequence Diagram of {TITLE}", width=6.04)

    section_heading(document, "3.2\tPermanent Data Storage")
    body(document, "The system uses SQLite for entities that require relationships, constraints, indexes, or private history. The database file is opened with foreign-key enforcement, and FTS5 is rebuilt during initialization so an older local database cannot silently omit new content from search. Schema migration code adds later animal and user columns without deleting existing rows. Ecosystem and protected-area catalogues remain local JSON snapshots because their assessment structures contain many source-specific fields and are primarily read-only in the member interface.")
    body(document, "At the time of inspection the working SQLite database contains 182 animals or protected taxon records, 506 Myanmar aliases, three categories, two users, twenty saved chat messages, 3,766 unique embedding vectors, 5,933 passage mappings, and one cached query vector. These counts describe the inspected workspace and are not hard-coded requirements. The application reads current totals at runtime.")

    subsection_heading(document, "3.2.1\tData Description")
    body(document, "The following data dictionaries are derived from the current schema in db.py and wildlife.db. Primary keys uniquely identify rows, foreign keys enforce relationships where the data model permits them, and application-level stable keys are identified explicitly when a database foreign key would cross a JSON catalogue or a derived subject key.")

    users_rows = [
        ("id", "integer", "Primary Key", "Unique user identifier."),
        ("name", "text", "Not Null", "Display name shown in the signed-in interface."),
        ("email", "text", "Not Null, Unique", "Normalized account email used to sign in."),
        ("password_hash", "text", "Not Null", "Werkzeug password hash; plaintext is never stored."),
        ("created_at", "text", "Not Null", "UTC account-creation timestamp."),
        ("is_admin", "integer", "Not Null, Default: 0", "Role flag granting Content Studio access when set."),
    ]
    data_table(document, "Table 3.1 User Table", users_rows)

    category_rows = [
        ("id", "integer", "Primary Key", "Unique category identifier."),
        ("name_my", "text", "Not Null", "Myanmar category label."),
        ("name_en", "text", "Not Null", "English category label."),
        ("description_my", "text", "Optional", "Myanmar category explanation."),
        ("icon", "text", "Optional", "Icon or symbol used by the interface."),
    ]
    data_table(document, "Table 3.2 Category Table", category_rows)

    animal_rows = [
        ("id", "integer", "Primary Key", "Unique animal or protected-taxon identifier."),
        ("name_my", "text", "Not Null", "Primary Myanmar display name."),
        ("name_en", "text", "Not Null", "Primary English display name."),
        ("scientific_name", "text", "Optional", "Scientific name or named taxonomic group."),
        ("category_id", "integer", "Foreign Key → categories.id", "Mammal, bird, or reptile category."),
        ("habitat_my", "text", "Optional", "Myanmar habitat description."),
        ("description_my", "text", "Optional", "Main Myanmar record description."),
        ("diet_my", "text", "Optional", "Myanmar diet description."),
        ("conservation_status", "text", "Optional", "National or legacy conservation code."),
        ("image_url", "text", "Optional", "Verified local cover photograph path."),
        ("fun_facts_my", "text", "Optional", "Additional Myanmar facts."),
        ("size_my", "text", "Optional", "Body-size information."),
        ("lifespan_my", "text", "Optional", "Lifespan information."),
        ("region_my", "text", "Optional", "Myanmar occurrence or region information."),
        ("source_record_id", "text", "Optional, Unique when present", "Stable identifier of the source record."),
        ("source_url", "text", "Optional", "Principal published source URL."),
        ("source_name", "text", "Optional", "Name of the principal source."),
        ("source_image_url", "text", "Optional", "Original source location of the lead image."),
        ("linked_name", "text", "Optional, Unique when present", "Stable portal-linked name used by importers."),
        ("other_name", "text", "Optional", "Alternative name supplied by the source."),
        ("species_order", "text", "Optional", "Taxonomic order."),
        ("species_family", "text", "Optional", "Taxonomic family."),
        ("global_iucn_status", "text", "Optional", "Global IUCN Red List category code."),
        ("population_trend", "text", "Optional", "Known global population direction."),
        ("myanmar_population", "text", "Optional", "Population information specific to Myanmar."),
        ("global_population", "text", "Optional", "Global population information."),
        ("portal_data", "text", "Optional", "Serialized source payload retained for provenance."),
        ("crawled_at", "text", "Optional", "Timestamp of portal acquisition."),
        ("male_image_url", "text", "Optional", "Verified local male photograph path."),
        ("female_image_url", "text", "Optional", "Verified local female photograph path."),
        ("source_urls", "text", "Optional", "Serialized list of supporting source URLs."),
        ("protection_status_my", "text", "Optional", "Myanmar legal-protection label."),
        ("protection_status_en", "text", "Optional", "English legal-protection label."),
        ("protection_law_my", "text", "Optional", "Myanmar legal schedule or notification wording."),
        ("protection_basis_my", "text", "Optional", "Why the record is protected directly or by group."),
        ("image_credit", "text", "Optional", "Photographer, source, licence, or group-image clarification."),
        ("taxon_rank", "text", "Optional; NULL means species", "Species, genus, family, or order rank."),
    ]
    data_table(document, "Table 3.3 Animal Table", animal_rows)

    alias_rows = [
        ("id", "integer", "Primary Key", "Unique alias identifier."),
        ("animal_id", "integer", "Foreign Key → animals.id, Not Null", "Animal record identified by the alias."),
        ("alias_my", "text", "Not Null", "Myanmar alternative name used in search and subject matching."),
    ]
    data_table(document, "Table 3.4 AnimalAlias Table", alias_rows)

    chat_rows = [
        ("id", "integer", "Primary Key", "Unique saved message identifier."),
        ("user_id", "integer", "Foreign Key → users.id, Not Null", "Owner of the conversation message."),
        ("animal_key", "text", "Not Null, Indexed with user", "Stable species-specific conversation key."),
        ("role", "text", "Not Null, Check: user|assistant", "Message author role."),
        ("content", "text", "Not Null", "User question or completed assistant response."),
        ("created_at", "text", "Not Null", "UTC saved timestamp."),
    ]
    data_table(document, "Table 3.5 ChatMessage Table", chat_rows)

    vector_rows = [
        ("model", "text", "Composite Primary Key", "Embedding model that produced the vector."),
        ("text_hash", "text", "Composite Primary Key", "SHA-256 hash of the exact passage text."),
        ("dims", "integer", "Not Null", "Number of vector dimensions."),
        ("vector", "blob", "Not Null", "Normalized float vector stored as a binary array."),
        ("created_at", "text", "Not Null", "UTC vector creation timestamp."),
    ]
    data_table(document, "Table 3.6 EmbeddingVector Table", vector_rows)

    passage_rows = [
        ("model", "text", "PK; FK → embedding_vectors.model", "Embedding model used for this mapping."),
        ("content_kind", "text", "PK; Check: animal|ecosystem|protected-area", "Catalogue containing the record."),
        ("record_key", "text", "Composite Primary Key", "Stable animal key or JSON catalogue slug."),
        ("position", "integer", "Composite Primary Key", "Ordered passage position in the record."),
        ("passage_label", "text", "Not Null", "Human-readable field or passage label."),
        ("text_hash", "text", "FK → embedding_vectors.text_hash, Not Null", "Vector hash for the passage text."),
        ("updated_at", "text", "Not Null", "UTC time when the mapping was refreshed."),
    ]
    data_table(document, "Table 3.7 EmbeddingPassage Table", passage_rows)

    query_rows = [
        ("model", "text", "Composite Primary Key", "Embedding model used for the cached question."),
        ("text_hash", "text", "Composite Primary Key", "Hash of the exact normalized question text."),
        ("dims", "integer", "Not Null", "Number of vector dimensions."),
        ("vector", "blob", "Not Null", "Normalized cached question vector."),
        ("created_at", "text", "Not Null", "UTC query-vector creation timestamp."),
    ]
    data_table(document, "Table 3.8 EmbeddingQuery Table", query_rows)

    fts_rows = [
        ("rowid", "integer", "Logical link → animals.id", "External-content row identifier."),
        ("name_my", "text", "Indexed", "Myanmar name searchable through FTS5."),
        ("name_en", "text", "Indexed", "English name searchable through FTS5."),
        ("scientific_name", "text", "Indexed", "Scientific name searchable through FTS5."),
        ("description_my", "text", "Indexed", "Myanmar description searchable through FTS5."),
    ]
    data_table(document, "Table 3.9 AnimalsFTS Virtual Table", fts_rows)

    json_rows = [
        ("slug / id", "text", "Stable logical key", "Record key used by routes and embedding mappings."),
        ("name / name_mm", "text", "Present", "English and Myanmar display names."),
        ("description fields", "text", "Optional", "Bilingual descriptions, distributions, assessments, or purposes."),
        ("classification fields", "text/list", "Optional", "Biome, realm, designation, status, region, authority, or protection fields."),
        ("source_name / source_url", "text", "Present", "Published provenance for the local snapshot."),
        ("image_url / image_urls", "text/list", "Optional", "Verified local WebP paths; missing images remain explicit."),
        ("source_image_urls", "list", "Optional", "Original image locations retained for provenance."),
    ]
    data_table(document, "Table 3.10 Local JSON Catalogue Record", json_rows)

    subsection_heading(document, "3.2.2\tDatabase Diagram")
    body(document, "The logical database diagram emphasizes the relationships used by the application. One category has many animal records, and one animal can have many aliases. One user has many chat messages. The chat animal_key is an application-level association rather than a foreign key because it is a stable conversational subject key. FTS5 uses animals as external content and is maintained through triggers. Embedding vectors are deduplicated by model and text hash, while passage mappings associate the same vector with every record and position that publishes the text. JSON catalogue slugs participate in the passage mapping through logical record keys.")
    figure(document, ASSETS / "diagrams" / "database-diagram.png", f"Figure 3.4: Database Diagram of {TITLE}", width=6.04)

    section_heading(document, "3.3\tSummary")
    body(document, "This chapter translates the project requirements into a seven-layer architecture, role-based use cases, a guarded question-answering sequence, and a permanent storage model. The browser and Jinja presentation layer remain separate from Flask access control, Myanmar NLP, retrieval, external AI, and local storage. The sequence design ensures that authentication and deterministic gates precede generation. SQLite supplies constrained relational data, private history, FTS5, and vector storage, while local JSON catalogues and verified media support rich read-only environmental records. Together these components create an inspectable path from published wildlife evidence to a member-facing answer.")

    chapter(document, 4, "Implementation")
    body(document, "The implementation turns the preceding architecture into a working desktop web application. Flask creates the application object, registers the administrator blueprint, initializes SQLite on the first request, loads the signed-in account, and renders Jinja templates. CSS and JavaScript provide the catalogue layout, search interaction, chat workspace, streamed Markdown, history rail, and administration tools. The screenshots in this chapter were captured from the running project at a 1440 × 900 desktop viewport using a temporary copy of the database, so the visual evidence does not alter the developer's working data.")
    body(document, "The member interface is intentionally organised around published records rather than around the AI model. The navigation first exposes species, ecosystems, and protected areas. A member reads a record before opening the chat that belongs to that record. Content Studio is visible only to an authorised administrator. The implementation descriptions below follow the normal journey from public entry to account access, browsing, grounded conversation, and administration.")

    section_heading(document, "4.1\tSystem Implementation")

    subsection_heading(document, "4.1.1\tPublic Landing Interface")
    body(document, "The landing page is the only public content page. It explains the Myanmar wildlife focus, provides registration and sign-in actions, and uses representative local photographs for a species, ecosystem, and protected area. The summary figures are read from the current catalogues. In the captured system they show 182 animal or taxon records, 64 ecosystems, and 73 protected areas. A signed-in member who opens the root URL is redirected to the encyclopedia rather than shown the marketing page again.")
    figure(document, ASSETS / "screenshots" / "01-landing-desktop.png", "Figure 4.1: Public landing page in desktop view")

    subsection_heading(document, "4.1.2\tRegistration Interface")
    body(document, "A new user supplies a name, email address, password, and password confirmation. Server-side validation checks required values, password length, matching confirmation, and unique email. The password is converted to a secure hash before the user row is stored. Successful registration redirects to sign-in with a confirmation notice; it does not automatically make the first user an administrator.")
    figure(document, ASSETS / "screenshots" / "02-register-desktop.png", "Figure 4.2: Registration page in desktop view")

    subsection_heading(document, "4.1.3\tSign-In Interface")
    body(document, "The sign-in page verifies the email and password against the stored hash and writes the user identifier to the signed session. When an anonymous member was redirected here from a protected page, a validated same-site next path returns the member to that page. External or protocol-relative destinations are discarded. The header keeps registration and login visually separate so the account flow is clear.")
    figure(document, ASSETS / "screenshots" / "03-login-desktop.png", "Figure 4.3: Sign-in page in desktop view")

    subsection_heading(document, "4.1.4\tSpecies Dashboard")
    body(document, "After authentication, the species dashboard presents the live category counts and a curated set of prominent Myanmar animals. Each card shows a verified photograph or an honest missing-photo state, IUCN status with an expanded meaning, legal-protection badge when applicable, Myanmar and scientific identity, and a short source description. Category links and the search field provide alternative navigation. The lower page continues to ecosystem and protected-area discovery and offers compact links to the species-specific chats.")
    figure(document, ASSETS / "screenshots" / "04-species-home-desktop.png", "Figure 4.4: Signed-in species dashboard in desktop view")

    subsection_heading(document, "4.1.5\tSpecies Detail Interface")
    body(document, "The species detail page is the evidence page behind the conversational experience. It combines identity, taxonomy, local and global conservation information, legal protection, habitat, population, sources, image credit, and available photographs. Suggested questions are generated from fields the record actually contains. The chat action always opens the conversation for this exact animal identifier, which prevents a generic assistant from losing the active subject.")
    figure(document, ASSETS / "screenshots" / "05-animal-detail-desktop.png", "Figure 4.5: Species detail page for the Asian Elephant")

    subsection_heading(document, "4.1.6\tSpecies-Specific Chat Workspace")
    body(document, "The chat page is a full-height workspace. A rail lists recent species conversations for the signed-in member, while the central stage names the current species, provides record-aware suggestions, shows messages, and keeps the composer available. Each history row and the current conversation have scoped delete controls. The browser posts a message to the species endpoint and renders server-sent answer deltas, Markdown, source groups, redirect actions, or local images. The rail becomes a drawer at narrower widths, but this report uses the requested desktop view.")
    figure(document, ASSETS / "screenshots" / "06-species-chat-desktop.png", "Figure 4.6: Species-specific AI chat workspace in desktop view")

    subsection_heading(document, "4.1.7\tMyanmar-Aware Search")
    body(document, "The search interface accepts Myanmar, English, scientific, alias, and description terms. The backend normalizes invisible characters and mark order, scores exact, phonetic, fuzzy, and FTS5 matches, and limits results below the relevance threshold. The same search endpoint supports interactive suggestions in the header, but the full results page can display a larger set with match explanations. Literal treatment of FTS syntax prevents a user string from becoming an unintended query expression.")
    figure(document, ASSETS / "screenshots" / "09-search-desktop.png", "Figure 4.7: Myanmar species search results in desktop view")

    subsection_heading(document, "4.1.8\tEcosystem Catalogue")
    body(document, "The ecosystem catalogue reads the 64-record local snapshot and offers text, biome, and assessment filters with pagination. Cards combine bilingual names, classification, conservation assessment, and locally verified images. The detail route presents distribution, environment, characteristic species, threats, criteria, and original source information when present. Ecosystems are reference records and deliberately do not have their own chatbot; a species chat that detects an ecosystem name hands the user to the ecosystem page.")
    figure(document, ASSETS / "screenshots" / "07-ecosystems-desktop.png", "Figure 4.8: Ecosystem catalogue in desktop view")

    subsection_heading(document, "4.1.9\tProtected-Area Catalogue")
    body(document, "The protected-area catalogue contains 73 local records with filters for designation, status, and region. A detail page can show governance, management authority, boundary, protection level, purpose, area, establishment, biodiversity, threats, tourism, photographs, and the portal source. Like ecosystems, protected areas remain contextual reference pages rather than conversational subjects. This keeps the chat contract narrow and prevents a species prompt from pretending to contain a protected-area dossier.")
    figure(document, ASSETS / "screenshots" / "08-protected-areas-desktop.png", "Figure 4.9: Protected-area catalogue in desktop view")

    subsection_heading(document, "4.1.10\tContent Studio")
    body(document, "Content Studio is an administrator-only publishing surface. Its dashboard summarises species, ecosystems, and protected areas, then shows cards for the selected catalogue with live-page, edit, and delete actions. Animal and collection forms validate textual fields, sources, and image files. Uploads are checked against the configured request limit and recognised image signatures. CSRF tokens protect state-changing forms, and publishing a changed record can immediately refresh its passage index while an authorised administrator is online.")
    figure(document, ASSETS / "screenshots" / "10-content-studio-desktop.png", "Figure 4.10: Content Studio administration dashboard")

    subsection_heading(document, "4.1.11\tRetrieval Pipeline Console")
    body(document, "The retrieval console makes the RAG workflow visible. It reports the twelve-stage pipeline, vector-store size, record and passage coverage, maximum selected passages, embedding model, generator, active threshold, and record-specific probes. A question run exposes normalization, Myanmar segmentation, stopword filtering, subject recognition, scope and photograph gates, passage inventory, intent tags, lexical and cosine scores, selected evidence, prompt context, and the final branch. Deterministic branches display the exact member response; model generation is an explicit action so opening the console does not silently call the external generator.")
    figure(document, ASSETS / "screenshots" / "11-rag-pipeline-desktop.png", "Figure 4.11: Administrator retrieval-pipeline console")

    section_heading(document, "4.2\tInstallation Manual")
    body(document, "A development installation requires Python, the repository files, and network access only for operations that contact Groq, OpenAI, or live source sites. The checked-in SQLite database, JSON snapshots, and local photographs allow the encyclopedia and most retrieval behavior to run locally. Flask 3.1 supports Python 3.9 and newer [4]. The exact commands below are written for macOS or Linux first, followed by the Windows PowerShell variation where activation or file-copy syntax differs.")

    subsection_heading(document, "4.2.1\tOpen the Project Directory and Check Python")
    body(document, "Open a terminal and change to the Flask application directory. Confirm that the selected Python is available before creating the environment.")
    command_block(document, "cd /Users/myominko/Developer/animal/myanmar_wildlife\npython3 --version")
    body(document, "On Windows PowerShell, change to the cloned project path and use the Python launcher if it is installed.")
    command_block(document, "Set-Location C:\\path\\to\\animal\\myanmar_wildlife\npy -3 --version")

    subsection_heading(document, "4.2.2\tCreate and Activate a Virtual Environment")
    body(document, "A project-specific virtual environment prevents its Flask, request, fuzzy-search, dotenv, and image packages from changing other Python projects.")
    command_block(document, "python3 -m venv venv\nsource venv/bin/activate")
    body(document, "Windows PowerShell uses the Scripts directory inside the environment.")
    command_block(document, "py -3 -m venv venv\n.\\venv\\Scripts\\Activate.ps1")

    subsection_heading(document, "4.2.3\tInstall the Python Requirements")
    body(document, "Upgrade pip inside the activated environment, then install the versions listed in requirements.txt. The project directly lists Flask, requests, RapidFuzz, Werkzeug, python-dotenv, and Pillow; their transitive dependencies are installed automatically.")
    command_block(document, "python -m pip install --upgrade pip\npython -m pip install -r requirements.txt")

    subsection_heading(document, "4.2.4\tCreate the Environment Configuration")
    body(document, "Copy the provided example and edit the new .env file. SECRET_KEY must be replaced for any shared deployment. GROQ_API_KEY enables generated member answers. OPENAI_API_KEY is optional and is used only by administrator embedding operations. ADMIN_EMAILS may contain comma-separated accounts allowed to enter Content Studio. Never place either API key in templates or static JavaScript.")
    command_block(document, "cp .env.example .env\n# Edit .env with a text editor, for example:\nnano .env")
    body(document, "The equivalent PowerShell copy command is:")
    command_block(document, "Copy-Item .env.example .env\nnotepad .env")
    body(document, "A minimal local configuration can retain APP_HOST=127.0.0.1 and APP_PORT=2000. If no Groq key is configured, catalogue and deterministic retrieval branches still work, but an answerable generation request reports that the language model is unavailable.")

    subsection_heading(document, "4.2.5\tUse the Included Data or Rebuild from Local Caches")
    body(document, "The repository includes wildlife.db, ecosystem and protected-area JSON snapshots, and local image assets. No import is required for the ordinary demonstration. To rebuild the species database from the saved crawls without fetching the source pages again, run the three cache commands in order. The first command creates a timestamped backup before replacing generated portal data; the legal importers add or refresh the protected records they own.")
    command_block(document, "python crawler.py --from-cache\npython protected_mammals.py --from-cache\npython protected_birds.py --from-cache")
    body(document, "Running the importers without --from-cache performs live source work and can download photographs. That operation depends on network availability and should be reviewed before use in a controlled demonstration.")

    subsection_heading(document, "4.2.6\tRun the Development Server")
    body(document, "Start the Flask application from the myanmar_wildlife directory. The configured development address is http://127.0.0.1:2000. The built-in server is appropriate for local testing only; Flask's documentation recommends a dedicated WSGI server or hosting platform for production [4].")
    command_block(document, "python app.py\n# Open http://127.0.0.1:2000 in a desktop browser")

    subsection_heading(document, "4.2.7\tCreate an Administrator")
    body(document, "Register the account through the interface, then grant its email explicit Content Studio access with the Flask CLI. This avoids making the first registrant an administrator by accident. The same CLI supplies a revoke command.")
    command_block(document, "flask --app app promote-admin admin@example.com\n# To remove the role later:\nflask --app app revoke-admin admin@example.com")

    subsection_heading(document, "4.2.8\tBuild the Optional Embedding Index")
    body(document, "Set OPENAI_API_KEY and keep EMBED_ENABLED=1 before running the full passage-index command. It indexes animal, ecosystem, and protected-area passages with the configured model, reuses vectors whose text hash already exists, reports coverage, and prunes stale mappings only after a complete run. Member chat never performs this online indexing operation.")
    command_block(document, "flask --app app embed-warm\n# Development-only partial run:\nflask --app app embed-warm --limit 10")

    subsection_heading(document, "4.2.9\tRun Automated Verification")
    body(document, "The repository uses Python unittest. The full suite was executed while preparing this report and all 102 tests passed. The suite works on a temporary database copy and covers the catalogue, protected schedules, photographs, authentication, redirects, private history, Myanmar search, deterministic RAG gates, semantic recall, prompt boundaries, streaming, Content Studio, vector storage, and the live retrieval trace.")
    command_block(document, "python -m unittest discover -s tests -v")

    subsection_heading(document, "4.2.10\tOperational Notes")
    bullet(document, "Keep .env out of version control and rotate any key that has been exposed.")
    bullet(document, "Use a production WSGI server, HTTPS, secure session settings, backups, request logging, and an external reverse proxy before public deployment.")
    bullet(document, "Re-run the test suite after changing retrieval thresholds, Myanmar normalization rules, source importers, authentication, or the database schema.")
    bullet(document, "Use the retrieval console to compare lexical score, cosine score, selected passages, and the final prompt before blaming the generator for a weak answer.")
    bullet(document, "Preserve image credits and source URLs when editing a record; a visually attractive photograph is not acceptable if its subject or licence is uncertain.")

    section_heading(document, "4.3\tSummary")
    body(document, "This chapter demonstrates the implemented public, member, and administrator workflows with real desktop captures. The application provides a live-count landing page, validated account forms, species discovery, detailed evidence pages, species-specific chat, Myanmar-aware search, ecosystem and protected-area catalogues, Content Studio, and an inspectable retrieval console. The installation manual provides environment, dependency, configuration, cached-data, server, role, embedding, and test commands. The successful 102-test run confirms the current repository behavior across its principal data, security, search, RAG, media, history, and administration boundaries.")

    chapter(document, 5, "Conclusion")
    body(document, f"{TITLE} demonstrates that a useful wildlife chatbot is primarily a data and retrieval system, not merely a language-model interface. The completed application connects curated Myanmar wildlife records, environmental catalogues, legal protection, verified photographs, search, account access, deterministic safety branches, hybrid retrieval, streamed generation, private history, publishing, and diagnostics in one coherent desktop web experience.")
    body(document, "The implementation keeps the selected record at the centre of every conversational decision. A member cannot open a global chat; the page and endpoint name one species. A question about another record becomes a navigation hand-off. A photograph request reads local media metadata. A question with no supported evidence is refused before generation. Only a matched question reaches the model, and the prompt contains the retrieved passages rather than the complete catalogue. These boundaries make the system more predictable and easier to explain.")
    body(document, "The project also shows that Myanmar-language access requires deliberate engineering. Unicode normalization, mark-order handling, aliases, literal subject rules, head-noun ambiguity checks, phonetic comparison, fuzzy scoring, FTS5, syllable units, and intent-aware passages cooperate across catalogue search and chat retrieval. The same care appears in the distinction between conservation risk and legal protection, between a species and a protected taxonomic group, and between an actual photograph and an unavailable or example image.")
    body(document, "From an implementation perspective, Flask and Jinja provide a compact server-rendered application; SQLite provides relational data, FTS5, history, and vector storage; local JSON and WebP assets create complete environmental snapshots; Groq supplies streamed generation; and OpenAI embeddings provide administrator-prepared semantic recall. Content Studio and the retrieval console close the loop by allowing records to be published and retrieval decisions to be examined using the same application functions that serve members.")

    section_heading(document, "5.1\tLimitations")
    body(document, "The current system is a strong prototype, but its correctness is bounded by the quality and freshness of its source records. A grounded answer can still be incomplete or phrase uncertain evidence too strongly if the retrieved passage itself is incomplete. Retrieval limits the model's available context; it does not mathematically prove every generated sentence. Users should therefore treat the displayed sources and original record as the verification authority.")
    bullet(document, "External generation dependency: answerable member questions require a configured Groq key and service availability. Deterministic branches and local catalogues remain available without it, but matched natural-language answers do not.")
    bullet(document, "Prepared semantic recall: member chat never creates an online query embedding. An unseen question therefore uses lexical retrieval unless the exact question vector was previously cached by an administrator.")
    bullet(document, "Threshold calibration: EMBED_THRESHOLD and lexical scoring are application settings. A threshold suitable for the present Myanmar corpus may not transfer to a much larger or differently written collection.")
    bullet(document, "Snapshot freshness: ecosystem, protected-area, legal, and imported species data are local snapshots. They do not automatically reflect a source-site correction or a later conservation assessment until an authorised refresh is performed.")
    bullet(document, "Bilingual coverage: many fields have Myanmar text, but some imported descriptions fall back to English when no appropriate Myanmar source exists. Translation quality is therefore uneven across the catalogue.")
    bullet(document, "Source-level citations: the interface groups evidence sources, but it does not yet provide a sentence-to-passage citation marker for every generated claim.")
    bullet(document, "Single-file deployment: SQLite is appropriate for the current local application and test workload, but a high-concurrency multi-host public service would require a carefully managed database architecture.")
    bullet(document, "Basic account lifecycle: registration, sign-in, sign-out, and administrator promotion are implemented, while email verification, password reset, account deletion, multi-factor authentication, and fine-grained editorial roles are not yet complete product features.")
    bullet(document, "Production operations: the repository starts Flask's development server. A public deployment still needs a production WSGI server, HTTPS, secure cookie policy, secrets management, monitoring, backups, rate limits, and incident procedures.")
    bullet(document, "Evaluation scope: the 102 automated tests verify defined behavior and important edge cases, but they are not a human evaluation of every Myanmar answer, every imported fact, or every photograph's biological identity.")
    body(document, "These limitations do not invalidate the design; they define where the current guarantee ends. The application guarantees controlled routing, retrieval inputs, source metadata, deterministic branches, account scope, and tested software behavior. It does not guarantee that an external source is timeless or that a probabilistic generator is infallible.")

    section_heading(document, "5.2\tFuture Work")
    body(document, "Future development should strengthen evidence quality before expanding answer freedom. The highest-value work is systematic evaluation: assemble a reviewed Myanmar question set for habitat, diet, population, threats, law, photographs, sources, and deliberately unsupported topics; measure retrieval precision and recall; and tune thresholds against false answers rather than against visually pleasing demos.")
    bullet(document, "Add passage-level citation anchors so each answer sentence can link to the exact selected field, source URL, and local record section that supports it.")
    bullet(document, "Create a curator review workflow for generated answers, disputed facts, image identity, licence changes, and stale source records without converting user feedback directly into trusted content.")
    bullet(document, "Schedule safe source-difference reports that show additions, deletions, changed assessments, and image changes before an administrator accepts a new snapshot.")
    bullet(document, "Build a Myanmar retrieval benchmark containing canonical text, common spelling variants, invisible characters, mark-order variants, ambiguous head nouns, English and scientific names, and adversarial off-topic questions.")
    bullet(document, "Evaluate multilingual or locally hosted embedding models so arbitrary member questions can receive semantic comparison without sending new text to an external embedding service.")
    bullet(document, "Add richer geographic relationships among species, ecosystems, protected areas, and states or regions, including map layers whose sources and update dates are explicit.")
    bullet(document, "Introduce editorial roles such as viewer, contributor, reviewer, and publisher; add audit logs for record changes, photograph replacement, index operations, and administrator grants.")
    bullet(document, "Move high-scale deployments to a production database and object store while preserving the current stable record keys, source fields, and tested deterministic response contract.")
    bullet(document, "Add operational safeguards including rate limiting, structured logs, model-call metrics, latency and error dashboards, backup verification, and privacy-preserving retention rules.")
    bullet(document, "Improve accessibility through keyboard-flow review, screen-reader testing, focus visibility, colour-contrast checks, alternative text review, reduced-motion support, and bilingual error messages.")
    bullet(document, "Provide export and educational collection features only after their provenance model is defined, so exported facts and images retain the same sources, credits, and legal distinctions as the live record.")
    body(document, "The architecture already supports these directions because record identity, passage construction, generator calls, vector storage, catalogues, media, authentication, and administration are separated. Future components can therefore extend one boundary without redefining the evidence contract of every page.")

    section_heading(document, "5.3\tSummary")
    body(document, f"This report concludes that {TITLE} successfully implements a source-grounded, Myanmar-oriented wildlife encyclopedia with Retrieval-Augmented Generation and natural-language processing. The system publishes 182 animal or protected-taxon records, 64 ecosystems, and 73 protected areas in the inspected workspace; supports account-controlled browsing, Myanmar-aware search, species evidence, private chat history, verified local media, Content Studio, and a twelve-stage retrieval console; and stores relational, full-text, and vector data locally. Its deterministic subject, photograph, and scope gates ensure that the language model is used only after the application has established a valid evidence path.")
    body(document, "The delivered architecture, use-case, sequence, and database diagrams describe the implemented project rather than the unrelated examples supplied for visual format. The installation chapter provides runnable commands for dependencies, environment configuration, cached data, administration, embeddings, server startup, and tests. All 102 automated tests passed during report preparation. The remaining work is primarily stronger evaluation, fresher sources, finer citations, richer administration, and production operations. Those improvements can build on the present design without weakening its central principle: when the local wildlife record does not support an answer, the system must say so.")

    section_heading(document, "References")
    add_reference(document, 1, "Myanmar Wildlife Encyclopedia Project, “README, Flask application, retrieval modules, database schema, templates, static assets, and automated tests,” local repository, 2026.")
    add_reference(document, 2, "P. Lewis et al., “Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks,” arXiv:2005.11401, 2020. [Online]. Available: https://arxiv.org/abs/2005.11401. Accessed: Sep. 1, 2026.")
    add_reference(document, 3, "SQLite Consortium, “SQLite FTS5 Extension,” SQLite Documentation. [Online]. Available: https://www.sqlite.org/fts5.html. Accessed: Sep. 1, 2026.")
    add_reference(document, 4, "Pallets Projects, “Flask Documentation 3.1.x,” 2026. [Online]. Available: https://flask.palletsprojects.com/en/stable/. Accessed: Sep. 1, 2026.")
    add_reference(document, 5, "OpenAI, “Vector Embeddings,” OpenAI API Documentation. [Online]. Available: https://developers.openai.com/api/docs/guides/embeddings. Accessed: Sep. 1, 2026.")
    add_reference(document, 6, "Groq, “OpenAI Compatibility,” GroqDocs. [Online]. Available: https://console.groq.com/docs/openai. Accessed: Sep. 1, 2026.")
    add_reference(document, 7, "Myanmar Biodiversity Portal, “All Species,” 2026. [Online]. Available: https://www.myanmarbiodiversity.org/species. Accessed: Sep. 1, 2026.")
    add_reference(document, 8, "Myanmar Biodiversity Portal, “Ecosystems,” local curated snapshot and source records, 2026. [Online]. Available: https://www.myanmarbiodiversity.org/ecosystems. Accessed: Sep. 1, 2026.")
    add_reference(document, 9, "Myanmar Biodiversity Portal, “Protected Areas,” local curated snapshot and source records, 2026. [Online]. Available: https://www.myanmarbiodiversity.org/protected_areas. Accessed: Sep. 1, 2026.")
    add_reference(document, 10, "Forest Department, Ministry of Natural Resources and Environmental Conservation, Myanmar, “Protected Wildlife List, Notification No. 690/2020,” 2020; structured schedules retained in the project data directory.")
    add_reference(document, 11, "Python Software Foundation, “venv — Creation of Virtual Environments,” Python Documentation. [Online]. Available: https://docs.python.org/3/library/venv.html. Accessed: Sep. 1, 2026.")
    add_reference(document, 12, "Pallets Projects, “Werkzeug Security Helpers,” Werkzeug Documentation. [Online]. Available: https://werkzeug.palletsprojects.com/en/stable/utils/#module-werkzeug.security. Accessed: Sep. 1, 2026.")
    add_reference(document, 13, "Pallets Projects, “Jinja Documentation,” [Online]. Available: https://jinja.palletsprojects.com/. Accessed: Sep. 1, 2026.")
    add_reference(document, 14, "IUCN, “The IUCN Red List of Threatened Species,” [Online]. Available: https://www.iucnredlist.org/. Accessed: Sep. 1, 2026.")

    # CONTENT_INSERTION_POINT

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
