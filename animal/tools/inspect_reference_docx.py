from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


path = Path("/Users/myominko/Downloads/Final Body UniSync.docx")
doc = Document(path)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
section = doc.sections[0]
print(
    "section",
    section.page_width.inches,
    section.page_height.inches,
    section.top_margin.inches,
    section.bottom_margin.inches,
    section.left_margin.inches,
    section.right_margin.inches,
    section.header_distance.inches,
    section.footer_distance.inches,
)


def value_or_none(value, attr="pt"):
    if value is None:
        return None
    return round(getattr(value, attr), 3)


def describe_paragraph(index: int, para: Paragraph):
    text = " ".join(para.text.split())
    pf = para.paragraph_format
    run = next((r for r in para.runs if r.text.strip()), None)
    if run is None:
        run_info = None
    else:
        rpr = run._element.rPr
        fonts = rpr.rFonts if rpr is not None else None
        run_info = {
            "text": run.text[:45],
            "name": run.font.name,
            "ascii": fonts.get(qn("w:ascii")) if fonts is not None else None,
            "size": value_or_none(run.font.size),
            "bold": run.bold,
            "italic": run.italic,
        }
    print(
        index,
        repr(text[:120]),
        {
            "style": para.style.name,
            "align": para.alignment,
            "before": value_or_none(pf.space_before),
            "after": value_or_none(pf.space_after),
            "line": pf.line_spacing,
            "line_rule": pf.line_spacing_rule,
            "first": value_or_none(pf.first_line_indent, "inches"),
            "left": value_or_none(pf.left_indent, "inches"),
            "right": value_or_none(pf.right_indent, "inches"),
            "keep_next": pf.keep_with_next,
            "page_break": pf.page_break_before,
            "run": run_info,
        },
    )


interesting = []
for index, para in enumerate(doc.paragraphs):
    text = " ".join(para.text.split())
    if index < 70 or text.startswith(("CHAPTER", "Figure", "Table", "4.1", "4.2", "5.1")):
        interesting.append((index, para))

for index, para in interesting[:180]:
    describe_paragraph(index, para)

print("\nSTYLE COUNTS")
print(Counter(p.style.name for p in doc.paragraphs).most_common())
print("\nTABLE GRIDS")
for index, table in enumerate(doc.tables[:20], start=1):
    widths = [cell.width.inches if cell.width else None for cell in table.rows[0].cells]
    print(index, len(table.rows), len(table.columns), widths, table.style.name if table.style else None)

print("\nINLINE SHAPES")
for index, shape in enumerate(doc.inline_shapes, start=1):
    print(index, round(shape.width.inches, 3), round(shape.height.inches, 3))
